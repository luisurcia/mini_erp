"""Aplica RUT/dirección/región/teléfono desde la planilla de contactos de
Scoby (`datos a migrar/0926 - Base de datos Clientes Kombucha.xlsx`) a los
clientes cuya fila ya fue confirmada en #140 — mapeo fila-Excel → cliente
en un CSV (`hoja,nombre_excel,customer_id`). Ver #141 para el detalle
completo de cada regla de normalización.

Solo completa campos hoy vacíos (nunca pisa un dato ya cargado a mano,
salvo `--force`) y nunca inventa un formato: RUT/teléfono/región que no
calzan con lo esperado se dejan sin escribir y se reportan para revisión
manual.

Los campos de la hoja Empresa sin columna propia en `Customer` (Razón
Social, Tipo de local, Contacto, Estado del local) quedan fuera de este
alcance — #140 todavía no confirmó si se quieren guardar en `notes`.

Corre desde la CLI:
`flask update-customer-contacts <excel.xlsx> <mapeo.csv> [--dry-run] [--force]`.
Kept out of any blueprint so it can be unit-tested on its own, same as
`app/migration.py`.
"""

from __future__ import annotations

import csv
import re
import unicodedata
from dataclasses import dataclass, field

from openpyxl import load_workbook

from app.exceptions import MiniErpError
from app.extensions import db
from app.models.customer import Customer
from app.repositories.customer_repository import CustomerRepository

SHEET_EMPRESA = "CLIENTE EMPRESA"
SHEET_PARTICULAR = "CLIENTE PARTICULAR"

# Comuna (sin tildes, minúscula) -> región. Acotada a las comunas que
# aparecen en esta planilla (ver #141) — no es un catastro nacional.
_COMUNA_REGION: dict[str, str] = {
    # Región Metropolitana
    "santiago": "Región Metropolitana",
    "santiago centro": "Región Metropolitana",
    "independencia": "Región Metropolitana",
    "la cisterna": "Región Metropolitana",
    "la florida": "Región Metropolitana",
    "la reina": "Región Metropolitana",
    "las condes": "Región Metropolitana",
    "nunoa": "Región Metropolitana",
    "penalolen": "Región Metropolitana",
    "providencia": "Región Metropolitana",
    "vitacura": "Región Metropolitana",
    "huechuraba": "Región Metropolitana",
    "san miguel": "Región Metropolitana",
    # Valparaíso
    "concon": "Valparaíso",
    "la calera": "Valparaíso",
    "la cruz": "Valparaíso",
    "quillota": "Valparaíso",
    "quilpue": "Valparaíso",
    "quilpue (los pinos)": "Valparaíso",
    "renaca": "Valparaíso",
    "vina del mar": "Valparaíso",
    "valparaiso": "Valparaíso",
    "nogales": "Valparaíso",
    "limache": "Valparaíso",
    "villa alemana": "Valparaíso",
    "laguna zapallar": "Valparaíso",
    "maitencillo": "Valparaíso",
}

# Código de región (hoja Empresa) -> nombre completo.
_REGION_CODES: dict[str, str] = {
    "RM": "Región Metropolitana",
    "V": "Valparaíso",
}


def _strip_accents(text: str) -> str:
    return "".join(
        c for c in unicodedata.normalize("NFD", text) if unicodedata.category(c) != "Mn"
    )


def _norm_key(raw) -> str:
    if raw is None:
        return ""
    return _strip_accents(" ".join(str(raw).split()).lower())


def region_for_comuna(raw) -> str | None:
    """Región inferida desde el nombre de comuna, o None si no está en la
    tabla — no se inventa una región para una comuna desconocida."""
    return _COMUNA_REGION.get(_norm_key(raw))


def normalize_phone(raw) -> tuple[str | None, str | None]:
    """Normaliza a `+56912345678`. Devuelve `(valor, None)` si calza, o
    `(None, motivo)` si no se pudo — nunca inventa un número."""
    if raw is None or not str(raw).strip():
        return None, None
    digits = re.sub(r"\D", "", str(raw))
    if len(digits) == 9 and digits.startswith("9"):
        digits = "56" + digits
    if re.fullmatch(r"569\d{8}", digits):
        return f"+{digits}", None
    return None, f"teléfono con formato irreconocible: {raw!r}"


def normalize_rut(raw) -> tuple[str | None, str | None]:
    """Normaliza a `NN.NNN.NNN-D`. Devuelve `(valor, None)` si calza, o
    `(None, motivo)` si es ambiguo/irreconocible — nunca inventa un
    dígito verificador (un número de 7-8 dígitos sin guion ni `K` final
    no se puede separar en número+dv de forma confiable)."""
    if raw is None or not str(raw).strip():
        return None, None
    compact = re.sub(r"[.\s]", "", str(raw)).upper()
    if "-" in compact:
        number, _, dv = compact.rpartition("-")
    elif compact.endswith("K"):
        number, dv = compact[:-1], "K"
    else:
        return None, f"RUT con formato irreconocible (sin separador): {raw!r}"
    if not re.fullmatch(r"\d{7,8}", number) or not re.fullmatch(r"[0-9K]", dv):
        return None, f"RUT con formato irreconocible: {raw!r}"
    formatted_number = f"{int(number):,}".replace(",", ".")
    return f"{formatted_number}-{dv}", None


@dataclass
class MapeoRow:
    hoja: str
    nombre_excel: str
    customer_id: int


def _read_mapeo(path: str) -> list[MapeoRow]:
    rows = []
    with open(path, newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        missing = {"hoja", "nombre_excel", "customer_id"} - set(reader.fieldnames or [])
        if missing:
            raise MiniErpError(f"Al mapeo.csv le faltan columnas: {', '.join(sorted(missing))}")
        for r in reader:
            hoja = r["hoja"].strip()
            if hoja not in (SHEET_EMPRESA, SHEET_PARTICULAR):
                raise MiniErpError(f"Hoja desconocida en el mapeo: {hoja!r}")
            rows.append(
                MapeoRow(
                    hoja=hoja,
                    nombre_excel=r["nombre_excel"].strip(),
                    customer_id=int(r["customer_id"]),
                )
            )
    return rows


def _index_sheet(ws, name_col: int) -> dict[str, list[tuple]]:
    """nombre normalizado -> lista de filas crudas de esa hoja (para
    detectar nombres ambiguos dentro de la misma hoja, p. ej. #275)."""
    index: dict[str, list[tuple]] = {}
    for row in ws.iter_rows(values_only=True):
        name = row[name_col] if name_col < len(row) else None
        if name is None or not str(name).strip():
            continue
        index.setdefault(_norm_key(name), []).append(row)
    return index


@dataclass
class ContactsUpdateResult:
    dry_run: bool
    source_name: str
    mapeo_name: str
    mapeo_rows: int = 0
    updated_customers: int = 0
    fields_updated: dict = field(default_factory=lambda: dict.fromkeys(
        ["rut", "shipping_street", "shipping_commune", "shipping_region", "phone"], 0
    ))
    skipped: list[tuple[str, str, str]] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    details: list[str] = field(default_factory=list)


def update_customer_contacts(
    xlsx_path: str, mapeo_path: str, *, dry_run: bool, force: bool = False
) -> ContactsUpdateResult:
    mapeo = _read_mapeo(mapeo_path)
    result = ContactsUpdateResult(
        dry_run=dry_run,
        source_name=xlsx_path.rsplit("/", 1)[-1],
        mapeo_name=mapeo_path.rsplit("/", 1)[-1],
        mapeo_rows=len(mapeo),
    )

    wb = load_workbook(xlsx_path, data_only=True, read_only=True)
    # Empresa: Nombre en columna 3 (0-based); Particular: Nombre Cliente en columna 1.
    empresa_index = _index_sheet(wb[SHEET_EMPRESA], name_col=3)
    particular_index = _index_sheet(wb[SHEET_PARTICULAR], name_col=1)
    wb.close()

    repo = CustomerRepository()

    try:
        for m in mapeo:
            customer = repo.get(m.customer_id)
            if customer is None:
                reason = f"cliente #{m.customer_id} no existe"
                result.skipped.append((m.hoja, m.nombre_excel, reason))
                continue

            index = empresa_index if m.hoja == SHEET_EMPRESA else particular_index
            candidates = index.get(_norm_key(m.nombre_excel), [])
            if not candidates:
                result.skipped.append(
                    (m.hoja, m.nombre_excel, "no se encontró esa fila en la hoja")
                )
                continue
            if len(candidates) > 1:
                reason = (
                    f"nombre ambiguo dentro de la hoja ({len(candidates)} filas) "
                    "— requiere revisión manual"
                )
                result.skipped.append((m.hoja, m.nombre_excel, reason))
                continue

            row = candidates[0]
            changed = _apply_row(customer, m.hoja, row, result, force=force)
            if changed:
                result.updated_customers += 1

        if dry_run:
            db.session.rollback()
        else:
            db.session.commit()
    except Exception:
        db.session.rollback()
        raise

    return result


def _blank(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _set_field(customer, field_name, new_value, result, *, force) -> bool:
    if new_value is None:
        return False
    current = getattr(customer, field_name)
    if not force and not _blank(current):
        return False
    setattr(customer, field_name, new_value)
    result.fields_updated[field_name] += 1
    return True


def _apply_row(
    customer: Customer, hoja: str, row: tuple, result: ContactsUpdateResult, *, force: bool
) -> bool:
    label = f"#{customer.id} {customer.name}"
    updated_fields = []

    if hoja == SHEET_EMPRESA:
        # (None, Tipo de local, Razón Social, Nombre, Rut, Direccion, REGION,
        #  Comuna, Teléfono, Contacto, Estado)
        rut_raw = row[4]
        direccion_raw, region_code, comuna_raw, phone_raw = row[5], row[6], row[7], row[8]
    else:
        # (None, Nombre Cliente, Dirección, Comuna, Teléfono)
        rut_raw = None
        direccion_raw, comuna_raw, phone_raw = row[2], row[3], row[4]

    rut_value, rut_err = normalize_rut(rut_raw)
    if rut_err:
        result.warnings.append(f"{label}: {rut_err}")
    elif rut_value is not None:
        conflict = Customer.query.filter(
            Customer.rut == rut_value, Customer.id != customer.id
        ).first()
        if conflict is not None:
            result.warnings.append(
                f"{label}: RUT {rut_value} ya está asignado a "
                f"#{conflict.id} {conflict.name} — no se escribió."
            )
        elif _set_field(customer, "rut", rut_value, result, force=force):
            updated_fields.append("rut")

    if direccion_raw is not None and str(direccion_raw).strip():
        if _set_field(customer, "shipping_street", str(direccion_raw).strip(), result, force=force):
            updated_fields.append("shipping_street")

    if comuna_raw is not None and str(comuna_raw).strip():
        if _set_field(customer, "shipping_commune", str(comuna_raw).strip(), result, force=force):
            updated_fields.append("shipping_commune")

    if hoja == SHEET_EMPRESA:
        region_name = _REGION_CODES.get(str(region_code).strip().upper()) if region_code else None
        inferred = region_for_comuna(comuna_raw)
        if region_name and inferred and region_name != inferred:
            result.warnings.append(
                f"{label}: código de región {region_code!r} ({region_name}) no calza con "
                f"la región inferida de la comuna {comuna_raw!r} ({inferred}) — "
                f"se guardó {region_name}, revisar."
            )
        region_value = region_name
    else:
        region_value = region_for_comuna(comuna_raw)
        if region_value is None and comuna_raw:
            result.warnings.append(
                f"{label}: comuna {comuna_raw!r} no está en la tabla de "
                "conversión — región sin completar."
            )

    if region_value and _set_field(customer, "shipping_region", region_value, result, force=force):
        updated_fields.append("shipping_region")

    phone_value, phone_err = normalize_phone(phone_raw)
    if phone_err:
        result.warnings.append(f"{label}: {phone_err}")
    elif phone_value and _set_field(customer, "phone", phone_value, result, force=force):
        updated_fields.append("phone")

    if updated_fields:
        result.details.append(f"{label}: {', '.join(updated_fields)}")
        return True
    return False


def build_report(result: ContactsUpdateResult) -> str:
    mode = "SIMULACIÓN (dry-run, no se escribió nada)" if result.dry_run else "CARGA REAL"
    lines = [
        "# Log de actualización de contactos de clientes — planilla Scoby (#141)",
        "",
        f"- **Modo:** {mode}",
        f"- **Archivo Excel:** `{result.source_name}`",
        f"- **Mapeo:** `{result.mapeo_name}` ({result.mapeo_rows} filas)",
        "",
        "## Resumen",
        "",
        "| Campo | Clientes actualizados |",
        "|---|---|",
    ]
    for f_name, count in result.fields_updated.items():
        lines.append(f"| {f_name} | {count} |")
    lines.append(f"| **Clientes con al menos un campo nuevo** | **{result.updated_customers}** |")

    lines += ["", f"## Detalle por cliente ({len(result.details)})", ""]
    lines += [f"- {d}" for d in result.details] or ["_Ninguno._"]

    lines += ["", f"## Advertencias ({len(result.warnings)})", ""]
    lines += [f"- {w}" for w in result.warnings] or ["_Ninguna._"]

    lines += ["", f"## Filas del mapeo omitidas ({len(result.skipped)})", ""]
    if result.skipped:
        lines += ["| Hoja | Nombre Excel | Motivo |", "|---|---|---|"]
        lines += [f"| {h} | {n} | {reason} |" for h, n, reason in result.skipped]
    else:
        lines.append("_Ninguna._")

    lines.append("")
    return "\n".join(lines)
